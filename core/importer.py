"""CSV / Excel 匯入與匯出（桌遊、會員）。

欄位名稱支援中英文別名，例如「遊戲名稱」「名稱」「name」都會對應到 name。
"""
from __future__ import annotations

import csv
import io
import math
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable

from .database import Database
from .models import BoardGame, Member, normalize_tags, TAG_LABELS

GAME_FIELD_ALIASES = {
    "name": ["name", "名稱", "遊戲名稱", "桌遊名稱", "品名", "商品名稱", "中文名稱"],
    "name_en": ["name_en", "英文名稱", "英文名", "english name"],
    "min_players": ["min_players", "最少人數", "最小人數", "人數下限"],
    "max_players": ["max_players", "最多人數", "最大人數", "人數上限"],
    "players": ["players", "人數"],                       # 例如 "2-5"
    "min_time": ["min_time", "最短時間", "時間下限"],
    "max_time": ["max_time", "最長時間", "時間上限"],
    "time": ["time", "時間", "遊戲時間", "時長"],            # 例如 "30-45"
    "weight": ["weight", "難度", "重度"],
    "min_age": ["min_age", "年齡", "適合年齡", "建議年齡"],
    "tags": ["tags", "標籤", "類型", "分類", "類別"],
    "description": ["description", "介紹", "簡介", "說明", "描述"],
    "pos_sku": ["pos_sku", "sku", "商品編號", "貨號", "條碼", "品號", "商品代碼"],
    "price": ["price", "售價", "價格", "定價"],
    "stock": ["stock", "庫存", "庫存量", "數量"],
    "rentable": ["rentable", "可租借", "店內遊玩", "可試玩"],
    "shelf": ["shelf", "櫃位", "位置", "架位"],
    "youtube_url": ["youtube_url", "youtube", "教學影片", "教學連結", "影片連結"],
    "image_url": ["image_url", "圖片", "圖片網址"],
    "active": ["active", "啟用", "上架"],
}

MEMBER_FIELD_ALIASES = {
    "member_no": ["member_no", "會員編號", "會員號碼", "卡號", "會員卡號", "編號"],
    "name": ["name", "姓名", "會員姓名", "名字"],
    "phone": ["phone", "手機", "電話", "手機號碼", "行動電話"],
    "level": ["level", "等級", "會員等級"],
    "points": ["points", "點數", "積分", "紅利"],
}

HISTORY_FIELD_ALIASES = {
    "member_no": ["member_no", "會員編號", "會員號碼", "卡號", "會員卡號"],
    "pos_sku": ["pos_sku", "sku", "商品編號", "貨號", "條碼", "品號", "商品代碼"],
    "game_name": ["game_name", "遊戲名稱", "桌遊名稱", "品名", "商品名稱"],
    "kind": ["kind", "類型", "交易類型", "種類", "銷售類型"],
    "occurred_at": ["occurred_at", "日期", "時間", "交易日期", "交易時間", "銷售日期", "租借日期"],
}

RENT_WORDS = {"rent", "rental", "租", "租借", "租用", "出租", "店內遊玩", "場租", "r"}
BUY_WORDS = {"buy", "sale", "sell", "purchase", "買", "購買", "銷售", "售出", "販售", "s", "b"}

TRUE_WORDS = {"1", "true", "yes", "y", "是", "v", "o", "可", "可以", "有"}


def _clean_num_text(v) -> str:
    return str(v).strip().replace(",", "").replace("$", "").replace("NT", "").replace("元", "")


def _to_int(v, default=0) -> int:
    try:
        f = float(_clean_num_text(v))
        return int(f) if math.isfinite(f) else default
    except (TypeError, ValueError, OverflowError):
        return default


def _to_float(v, default=0.0) -> float:
    try:
        f = float(_clean_num_text(v))
        return f if math.isfinite(f) else default
    except (TypeError, ValueError):
        return default


def _cell_text(v) -> str:
    """Excel 的整數儲存格可能讀成 1001.0，轉回 '1001'。"""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v if v is not None else "").strip()


def normalize_phone(v) -> str:
    p = _cell_text(v)
    for ch in "-() ":
        p = p.replace(ch, "")
    if p.startswith("+886"):
        p = "0" + p[4:]
    if len(p) == 9 and p.startswith("9") and p.isdigit():   # Excel 數字格式吃掉手機開頭的 0
        p = "0" + p
    return p


def _blank(v) -> bool:
    return v is None or str(v).strip() == ""


def _to_bool(v, default=True) -> bool:
    if v is None or str(v).strip() == "":
        return default
    return str(v).strip().lower() in TRUE_WORDS


def _parse_range(v) -> tuple[int, int] | None:
    """'2-5'、'2~5'、'2至5'、'4' → (2,5) 或 (4,4)。"""
    if v is None or isinstance(v, (datetime, date, time)):
        return None      # Excel 把「2-4」自動轉成日期時無法還原
    text = str(v).strip().replace("+", "").replace("人", "").replace("分鐘", "").replace("分", "")
    for sep in ["~", "～", "至", "到", "–", "—"]:
        text = text.replace(sep, "-")
    parts = [p for p in text.split("-") if p.strip()]
    nums = [_to_int(p, -1) for p in parts]
    nums = [n for n in nums if n >= 0]
    if not nums:
        return None
    return (min(nums), max(nums))


def _map_headers(headers: Iterable[str], aliases: dict) -> dict[str, str]:
    """回傳 {原始欄名: 內部欄名}。"""
    lookup = {}
    for field, names in aliases.items():
        for n in names:
            lookup[n.strip().lower()] = field
    mapping = {}
    for h in headers:
        if h is None:
            continue
        key = str(h).strip().lower()
        if key in lookup:
            mapping[h] = lookup[key]
    return mapping


class TableReader:
    """讀取 CSV（自動判斷 UTF-8 / Big5）或 Excel，輸出 list[dict]。"""

    @staticmethod
    def read(filename: str, data: bytes) -> list[dict]:
        ext = Path(filename).suffix.lower()
        if ext in (".xlsx", ".xlsm"):
            return TableReader._read_xlsx(data)
        if ext in (".csv", ".txt", ""):
            return TableReader._read_csv(data)
        if ext == ".xls":
            raise ValueError("不支援舊版 Excel .xls，請在 Excel 另存為 .xlsx 或 CSV")
        raise ValueError(f"不支援的檔案格式：{ext}（請使用 .csv 或 .xlsx）")

    @staticmethod
    def _read_csv(data: bytes) -> list[dict]:
        text = None
        for enc in ("utf-8-sig", "cp950", "big5", "utf-16"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("無法判斷 CSV 編碼，請另存為 UTF-8")
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel
        return [dict(r) for r in csv.DictReader(io.StringIO(text), dialect=dialect)]

    @staticmethod
    def _read_xlsx(data: bytes) -> list[dict]:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise ValueError(f"無法讀取 Excel 檔（檔案可能損毀或不是 .xlsx）：{e}") from e
        try:
            ws = wb.active or (wb.worksheets[0] if wb.worksheets else None)
            rows = list(ws.iter_rows(values_only=True)) if ws is not None else []
        finally:
            wb.close()
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        out = []
        for r in rows[1:]:
            if r is None or all(v is None or str(v).strip() == "" for v in r):
                continue
            out.append({headers[i]: r[i] for i in range(min(len(headers), len(r))) if headers[i]})
        return out


class ImportResult:
    def __init__(self):
        self.inserted = 0
        self.updated = 0
        self.skipped = 0
        self.errors: list[str] = []

    def summary(self) -> str:
        s = f"新增 {self.inserted} 筆、更新 {self.updated} 筆"
        if self.skipped:
            s += f"、略過 {self.skipped} 筆"
        return s


class GameImporter:
    def __init__(self, db: Database):
        self.db = db

    @staticmethod
    def provided_fields(raw: dict, mapping: dict[str, str]) -> set[str]:
        """這一列實際有填值的內部欄位（players / time 會展開成 min_/max_）。"""
        out = set()
        for k, v in raw.items():
            f = mapping.get(k)
            if f is None or _blank(v):
                continue
            if f in ("players", "time"):
                if _parse_range(v) is not None:
                    out |= {f"min_{f}", f"max_{f}"}
            else:
                out.add(f)
        return out

    def row_to_game(self, raw: dict, mapping: dict[str, str], source: str) -> BoardGame | None:
        d = {mapping[k]: v for k, v in raw.items() if k in mapping}
        name = _cell_text(d.get("name"))
        if not name:
            return None
        g = BoardGame(name=name, source=source)
        g.name_en = str(d.get("name_en") or "").strip()

        pr = _parse_range(d.get("players"))
        g.min_players = _to_int(d.get("min_players"), pr[0] if pr else 2)
        g.max_players = _to_int(d.get("max_players"), pr[1] if pr else max(g.min_players, 4))
        if g.max_players < g.min_players:
            g.max_players = g.min_players

        tr = _parse_range(d.get("time"))
        g.min_time = _to_int(d.get("min_time"), tr[0] if tr else 30)
        g.max_time = _to_int(d.get("max_time"), tr[1] if tr else max(g.min_time, 60))
        if g.max_time < g.min_time:
            g.max_time = g.min_time

        g.weight = min(5.0, max(1.0, _to_float(d.get("weight"), 2.0)))
        g.min_age = _to_int(str(d.get("min_age") or "8").replace("+", "").replace("歲", ""), 8)
        g.tags = normalize_tags(d.get("tags"))
        g.description = str(d.get("description") or "").strip()
        g.pos_sku = _cell_text(d.get("pos_sku"))
        g.price = _to_float(d.get("price"), 0)
        g.stock = _to_int(d.get("stock"), 0)
        g.rentable = _to_bool(d.get("rentable"), True)
        g.shelf = str(d.get("shelf") or "").strip()
        yt = str(d.get("youtube_url") or "").strip()
        if yt:
            g.youtube_url = yt
            g.youtube_is_search = "results?search_query" in yt
        g.image_url = str(d.get("image_url") or "").strip()
        g.active = _to_bool(d.get("active"), True)
        return g

    def import_rows(self, rows: list[dict], source: str = "import") -> ImportResult:
        res = ImportResult()
        if not rows:
            res.errors.append("檔案沒有資料")
            return res
        mapping = _map_headers(rows[0].keys(), GAME_FIELD_ALIASES)
        if "name" not in mapping.values():
            res.errors.append("找不到「遊戲名稱 / name」欄位，請確認標題列")
            return res
        for i, raw in enumerate(rows, start=2):
            try:
                g = self.row_to_game(raw, mapping, source)
                if g is None:
                    res.skipped += 1
                    continue
                if self.db.upsert_game(g, provided=self.provided_fields(raw, mapping)) == "inserted":
                    res.inserted += 1
                else:
                    res.updated += 1
            except Exception as e:  # 單筆錯誤不中斷整批
                res.errors.append(f"第 {i} 列：{e}")
                res.skipped += 1
        return res

    def import_file(self, filename: str, data: bytes) -> ImportResult:
        return self.import_rows(TableReader.read(filename, data))

    def export_csv(self) -> bytes:
        cols = ["name", "name_en", "min_players", "max_players", "min_time", "max_time",
                "weight", "min_age", "tags", "description", "pos_sku", "price", "stock",
                "rentable", "shelf", "youtube_url", "image_url", "active"]
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(cols)
        for g in self.db.list_games():
            d = g.to_dict()
            d["tags"] = "、".join(TAG_LABELS.get(t, t) for t in g.tags)
            d["rentable"] = int(g.rentable)
            d["active"] = int(g.active)
            w.writerow([d[c] for c in cols])
        return buf.getvalue().encode("utf-8-sig")   # 加 BOM，Excel 開啟才不會亂碼


def _history_time(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return _cell_text(v)[:30]


def _history_kind(v, default: str = "buy") -> str | None:
    """銷售/租借欄 → 'buy' / 'rent'；空白用預設；看不懂回傳 None。"""
    t = _cell_text(v).lower()
    if not t:
        return default
    if t in RENT_WORDS or "租" in t:
        return "rent"
    if t in BUY_WORDS or "售" in t or "買" in t:
        return "buy"
    return None


class HistoryImporter:
    """會員銷售／租借紀錄（CSV、Excel 或 POS 查詢結果）。依商品編號對應到桌遊。"""

    def __init__(self, db: Database):
        self.db = db

    def import_rows(self, rows: list[dict], source: str = "import", default_kind: str = "buy") -> ImportResult:
        res = ImportResult()
        if not rows:
            res.errors.append("檔案沒有資料")
            return res
        if default_kind not in ("buy", "rent"):
            default_kind = "buy"
        mapping = _map_headers(rows[0].keys(), HISTORY_FIELD_ALIASES)
        fields = set(mapping.values())
        if "member_no" not in fields:
            res.errors.append("找不到「會員編號 / member_no」欄位")
            return res
        if not fields & {"pos_sku", "game_name"}:
            res.errors.append("找不到「商品編號 / pos_sku」欄位")
            return res
        sku_map = self.db.sku_to_game_ids()
        name_cache: dict[str, int | None] = {}
        unmatched = 0
        batch: list[tuple] = []
        for i, raw in enumerate(rows, start=2):
            try:
                d = {mapping[k]: v for k, v in raw.items() if k in mapping}
                no = _cell_text(d.get("member_no"))
                sku = _cell_text(d.get("pos_sku"))
                gname = _cell_text(d.get("game_name"))
                if not no or not (sku or gname):
                    res.skipped += 1
                    continue
                kind = _history_kind(d.get("kind"), default_kind)
                if kind is None:
                    res.errors.append(f"第 {i} 列：看不懂交易類型「{d.get('kind')}」（請用 buy/rent 或 購買/租借）")
                    res.skipped += 1
                    continue
                gid = sku_map.get(sku) if sku else None
                if gid is None and gname:
                    if gname not in name_cache:
                        g = self.db.find_game_by_sku_or_name("", gname)
                        name_cache[gname] = g.id if g else None
                    gid = name_cache[gname]
                    if not sku and gid is not None:
                        sku = self.db.get_game(gid).pos_sku or f"#name:{gname}"
                    elif not sku:
                        sku = f"#name:{gname}"
                if gid is None:
                    unmatched += 1
                batch.append((no, sku, kind, _history_time(d.get("occurred_at")), source, gid))
            except Exception as e:  # 單筆錯誤不中斷整批
                res.errors.append(f"第 {i} 列：{e}")
                res.skipped += 1
        # 一次交易批次寫入（逐筆 commit 時 2 萬筆要 40 秒，POS 同步每次都會重抓全部紀錄）
        for start in range(0, len(batch), 5000):
            chunk = batch[start:start + 5000]
            n = self.db.add_history_many(chunk)
            res.inserted += n
            res.skipped += len(chunk) - n      # 已匯入過（重複）
        if unmatched:
            res.errors.append(f"有 {unmatched} 筆找不到對應的桌遊（商品編號不在桌遊資料庫）；"
                              "之後補上該桌遊的商品編號就會自動對上")
        return res

    def import_file(self, filename: str, data: bytes) -> ImportResult:
        return self.import_rows(TableReader.read(filename, data))


class MemberImporter:
    def __init__(self, db: Database):
        self.db = db

    def import_rows(self, rows: list[dict], source: str = "import") -> ImportResult:
        res = ImportResult()
        if not rows:
            res.errors.append("檔案沒有資料")
            return res
        mapping = _map_headers(rows[0].keys(), MEMBER_FIELD_ALIASES)
        if "member_no" not in mapping.values():
            res.errors.append("找不到「會員編號 / member_no」欄位")
            return res
        cols = set(mapping.values())
        for i, raw in enumerate(rows, start=2):
            try:
                d = {mapping[k]: v for k, v in raw.items() if k in mapping}
                no = _cell_text(d.get("member_no"))
                if not no:
                    res.skipped += 1
                    continue
                m = Member(member_no=no,
                           name=_cell_text(d.get("name")),
                           phone=normalize_phone(d.get("phone")),
                           level=_cell_text(d.get("level")),
                           points=_to_int(d.get("points"), 0),
                           source=source)
                if self.db.upsert_member(m, keep=({"name", "phone", "level", "points"} - cols)) == "inserted":
                    res.inserted += 1
                else:
                    res.updated += 1
            except Exception as e:  # 單筆錯誤不中斷整批
                res.errors.append(f"第 {i} 列：{e}")
                res.skipped += 1
        return res

    def import_file(self, filename: str, data: bytes) -> ImportResult:
        return self.import_rows(TableReader.read(filename, data))
